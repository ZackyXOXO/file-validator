import os  
import asyncio  
import logging  
import tempfile  

import pandas as pd  

from dotenv import load_dotenv  
from telegram import Update  
from telegram.request import HTTPXRequest  
from telegram.ext import (  
    Application,  
    CommandHandler,  
    MessageHandler,  
    filters,  
    ContextTypes  
)  

# ======================================  
# CONFIGURATION  
# ======================================  

load_dotenv()  

TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")  

if not TOKEN:  
    raise ValueError(  
        "TELEGRAM_BOT_TOKEN tidak ditemukan di file .env"  
    )  

logging.basicConfig(  
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",  
    level=logging.INFO  
)  

BASE_COLUMNS = [  
    "polygon",  
    "polygon_type",  
    "label"  
]  

UPDATE_COLUMNS = [  
    "polygon",  
    "polygon_type",  
    "label",  
    "label_new"  
]  

MAP_BORDER_FILE = "data/map_border.csv"  

MAP_BORDER_COLUMNS = [  
    "PROVINSI",  
    "KABUPATEN",  
    "KECAMATAN",  
    "DESA",  
    "AREA",  
    "ID_AREA"  
]  

TELEGRAM_TIMEOUT = 3600.0  
PROGRESS_INTERVAL = 300  

# ======================================  
# FORMAT TEXT  
# ======================================  

def format_title_case(value):  
    value = str(value).strip()  

    if not value:  
        return ""  

    return value[:1].upper() + value[1:].lower()  

# ======================================  
# LOAD MAP BORDER  
# ======================================  

def load_map_border():  
    if not os.path.exists(MAP_BORDER_FILE):  
        raise FileNotFoundError(  
            f"File Map Border tidak ditemukan: {MAP_BORDER_FILE}"  
        )  

    map_df = pd.read_csv(  
        MAP_BORDER_FILE,  
        dtype=str,  
        keep_default_na=False  
    )  

    map_df.columns = (  
        map_df.columns  
        .str.strip()  
        .str.upper()  
    )  

    missing_columns = [  
        col  
        for col in MAP_BORDER_COLUMNS  
        if col not in map_df.columns  
    ]  

    if missing_columns:  
        raise ValueError(  
            "Header Map Border tidak sesuai.\n"  
            f"Kolom kurang: {missing_columns}\n"  
            f"Header ditemukan: {list(map_df.columns)}"  
        )  

    for col in MAP_BORDER_COLUMNS:  
        map_df[col] = (  
            map_df[col]  
            .fillna("")  
            .astype(str)  
            .str.strip()  
            .str.upper()  
        )  

    return map_df[MAP_BORDER_COLUMNS].copy()  

# ======================================  
# VALIDATE MAP BORDER  
# ======================================  

def validate_map_border(row, map_df):  
    polygon = str(row["polygon"]).strip()  

    polygon_type = str(  
        row.get(  
            "_polygon_type_upper",  
            row["polygon_type"]  
        )  
    ).strip().upper()  
  
    if polygon_type == "GRID":  
        return None  

    if "_" not in polygon:  
        return (  
            f"Format polygon tidak sesuai untuk {polygon_type}. "  
            "Gunakan format KABUPATEN_NAMA"  
        )  

    kabupaten, wilayah = polygon.split("_", 1)  

    kabupaten = kabupaten.strip().upper()  
    wilayah = wilayah.strip().upper()  

    if not kabupaten or not wilayah:  
        return "Kabupaten atau wilayah polygon kosong"  

    if polygon_type == "DESA":  
        matched = map_df[  
            (map_df["KABUPATEN"] == kabupaten) &  
            (map_df["DESA"] == wilayah)  
        ]  

        if matched.empty:  
            return (  
                "Polygon DESA tidak ditemukan di Map Border: "  
                f"{polygon}"  
            )  

    elif polygon_type == "CITY":  
        matched = map_df[  
            (map_df["KABUPATEN"] == kabupaten) &  
            (map_df["AREA"] == wilayah)  
        ]  

        if matched.empty:  
            return (  
                "Polygon CITY tidak ditemukan di Map Border: "  
                f"{polygon}"  
            )  

    else:  
        return (  
            f"POLYGON_TYPE tidak valid: {polygon_type}. "  
            "Gunakan DESA, CITY, atau GRID"  
        )  

    return None  

# ======================================  
# MAP BORDER DETAIL RESULT  
# ======================================  

def create_map_border_result(df, map_df):  
    results = []  

    for _, row in df.iterrows():  
        polygon = str(row["polygon"]).strip()  

        polygon_type_upper = str(  
            row.get(  
                "_polygon_type_upper",  
                row["polygon_type"]  
            )  
        ).strip().upper()  

        polygon_type_display = format_title_case(  
            polygon_type_upper  
        )  

        result = {  
            "row_number": row["_original_row_number"],  
            "polygon": polygon,  
            "polygon_type": polygon_type_display,  
            "kabupaten": "",  
            "wilayah": "",  
            "status": "",  
            "id_area": "",  
            "keterangan": ""  
        }  

        if polygon_type_upper == "GRID":  
            result["status"] = "SKIPPED"  
            result["keterangan"] = (  
                "GRID tidak divalidasi ke Map Border"  
            )  
            results.append(result)  
            continue  

        if "_" not in polygon:  
            result["status"] = "INVALID FORMAT"  
            result["keterangan"] = (  
                "Format harus KABUPATEN_WILAYAH"  
            )  
            results.append(result)  
            continue  

        kabupaten, wilayah = polygon.split("_", 1)  

        kabupaten = kabupaten.strip().upper()  
        wilayah = wilayah.strip().upper()  

        result["kabupaten"] = kabupaten  
        result["wilayah"] = wilayah  

        if polygon_type_upper == "DESA":  
            matched = map_df[  
                (map_df["KABUPATEN"] == kabupaten) &  
                (map_df["DESA"] == wilayah)  
            ]  

        elif polygon_type_upper == "CITY":  
            matched = map_df[  
                (map_df["KABUPATEN"] == kabupaten) &  
                (map_df["AREA"] == wilayah)  
            ]  

        else:  
            result["status"] = "INVALID TYPE"  
            result["keterangan"] = (  
                "Gunakan DESA, CITY, atau GRID"  
            )  
            results.append(result)  
            continue  

        if matched.empty:  
            result["status"] = "NOT FOUND"  
            result["keterangan"] = (  
                "Polygon tidak ditemukan di Map Border"  
            )  
        else:  
            result["status"] = "MATCH"  
            result["id_area"] = matched.iloc[0]["ID_AREA"]  
            result["keterangan"] = (  
                "Polygon ditemukan di Map Border"  
            )  

        results.append(result)  

    return pd.DataFrame(results)  

# ======================================  
# VALIDATE CONTENT  
# ======================================  

def validate_content(df, expected_columns, map_df):  
    errors = []  

    for _, row in df.iterrows():  
        row_number = row["_original_row_number"]  
        polygon = str(row["polygon"])  

        for col in expected_columns:  
            value = row[col]  

            if pd.isna(value) or str(value).strip() == "":  
                errors.append({  
                    "row_number": row_number,  
                    "polygon": polygon,  
                    "error": f"{col.upper()} kosong"  
                })  

        if any(char in polygon for char in [",", ";", "|"]):  
            errors.append({  
                "row_number": row_number,  
                "polygon": polygon,  
                "error": (  
                    "Polygon mengandung karakter , ; atau |"  
                )  
            })  

        if polygon.strip():  
            map_border_error = validate_map_border(  
                row,  
                map_df  
            )  

            if map_border_error:  
                errors.append({  
                    "row_number": row_number,  
                    "polygon": polygon,  
                    "error": map_border_error  
                })  

    return errors  

# ======================================  
# CREATE EXCEL REPORT  
# ======================================  

def create_map_border_excel(  
    input_df,  
    map_border_result,  
    map_df,  
    output_file,  
    mode  
):  
    from openpyxl import load_workbook  
    from openpyxl.styles import Font, PatternFill, Alignment  
    from openpyxl.utils import get_column_letter  
    from openpyxl.worksheet.table import Table, TableStyleInfo  

    input_columns = [  
        "_original_row_number",  
        "polygon",  
        "polygon_type",  
        "label"  
    ]  

    if mode == "update":  
        input_columns.append("label_new")  

    input_result = input_df[input_columns].copy()  

    input_result = input_result.rename(columns={  
        "_original_row_number": "row_number"  
    })  

    input_result["kabupaten"] = ""  
    input_result["wilayah"] = ""  
    input_result["map_border_status"] = ""  
    input_result["id_area"] = ""  
    input_result["keterangan"] = ""  

    result_lookup = (  
        map_border_result  
        .set_index("row_number")  
        .to_dict("index")  
    )  

    for index, row in input_result.iterrows():  
        row_number = row["row_number"]  
        detail = result_lookup.get(row_number)  

        if detail:  
            input_result.at[index, "kabupaten"] = (  
                detail.get("kabupaten", "")  
            )  

            input_result.at[index, "wilayah"] = (  
                detail.get("wilayah", "")  
            )  

            input_result.at[index, "map_border_status"] = (  
                detail.get("status", "")  
            )  

            input_result.at[index, "id_area"] = (  
                detail.get("id_area", "")  
            )  

            input_result.at[index, "keterangan"] = (  
                detail.get("keterangan", "")  
            )  

    summary_polygon = (  
        input_result  
        .groupby(  
            ["polygon", "polygon_type"],  
            dropna=False  
        )  
        .size()  
        .reset_index(name="total_count")  
    )  

    summary_type = (  
        input_result  
        .groupby(  
            "polygon_type",  
            dropna=False  
        )  
        .size()  
        .reset_index(name="total_count")  
    )  

    summary_label = (  
        input_result  
        .groupby(  
            "label",  
            dropna=False  
        )  
        .size()  
        .reset_index(name="total_count")  
    )  

    with pd.ExcelWriter(  
        output_file,  
        engine="openpyxl"  
    ) as writer:  

        input_result.to_excel(  
            writer,  
            sheet_name="Validation Result",  
            index=False,  
            startrow=0  
        )  

        start_row = len(input_result) + 4  

        summary_polygon.to_excel(  
            writer,  
            sheet_name="Validation Result",  
            index=False,  
            startrow=start_row  
        )  

        start_row += len(summary_polygon) + 3  

        summary_type.to_excel(  
            writer,  
            sheet_name="Validation Result",  
            index=False,  
            startrow=start_row  
        )  

        start_row += len(summary_type) + 3  

        summary_label.to_excel(  
            writer,  
            sheet_name="Validation Result",  
            index=False,  
            startrow=start_row  
        )  

        map_df.to_excel(  
            writer,  
            sheet_name="Map Border",  
            index=False  
        )  

    workbook = load_workbook(output_file)  

    header_fill = PatternFill(  
        fill_type="solid",  
        fgColor="1F4E78"  
    )  

    header_font = Font(  
        bold=True,  
        color="FFFFFF"  
    )  

    for worksheet in workbook.worksheets:  
        worksheet.freeze_panes = "A2"  

        for cell in worksheet[1]:  
            cell.fill = header_fill  
            cell.font = header_font  
            cell.alignment = Alignment(  
                horizontal="center",  
                vertical="center"  
            )  

        for column_cells in worksheet.columns:  
            max_length = 0  

            column_letter = get_column_letter(  
                column_cells[0].column  
            )  

            for cell in column_cells:  
                try:  
                    max_length = max(  
                        max_length,  
                        len(str(cell.value))  
                    )  
                except Exception:  
                    pass  

            worksheet.column_dimensions[  
                column_letter  
            ].width = min(max_length + 2, 50)  

    validation_sheet = workbook["Validation Result"]  

    validation_headers = [  
        cell.value  
        for cell in validation_sheet[1]  
    ]  

    validation_last_row = len(input_result) + 1  
    validation_last_col = len(validation_headers)  

    if validation_last_row >= 2:  
        table_reference = (  
            f"A1:{get_column_letter(validation_last_col)}"  
            f"{validation_last_row}"  
        )  

        table = Table(  
            displayName="ValidationResultTable",  
            ref=table_reference  
        )  

        table.tableStyleInfo = TableStyleInfo(  
            name="TableStyleMedium2",  
            showFirstColumn=False,  
            showLastColumn=False,  
            showRowStripes=True,  
            showColumnStripes=False  
        )  

        validation_sheet.add_table(table)  

    map_sheet = workbook["Map Border"]  

    map_last_row = len(map_df) + 1  
    map_last_col = len(MAP_BORDER_COLUMNS)  

    if map_last_row >= 2:  
        table_reference = (  
            f"A1:{get_column_letter(map_last_col)}"  
            f"{map_last_row}"  
        )  

        table = Table(  
            displayName="MapBorderTable",  
            ref=table_reference  
        )  

        table.tableStyleInfo = TableStyleInfo(  
            name="TableStyleMedium4",  
            showFirstColumn=False,  
            showLastColumn=False,  
            showRowStripes=True,  
            showColumnStripes=False  
        )  

        map_sheet.add_table(table)  

    workbook.save(output_file)  

# ======================================  
# ADD VLOOKUP FORMULAS  
# ======================================  

def add_excel_vlookup_formulas(output_file):  
    from openpyxl import load_workbook  
    from openpyxl.styles import Font, PatternFill, Alignment  
    from openpyxl.utils import get_column_letter  

    workbook = load_workbook(output_file)  

    validation_sheet = workbook["Validation Result"]  
    map_sheet = workbook["Map Border"]  

    map_sheet["G1"] = "KEY_DESA"  
    map_sheet["H1"] = "ID_AREA_DESA"  
    map_sheet["I1"] = "KEY_CITY"  
    map_sheet["J1"] = "ID_AREA_CITY"  

    for row in range(2, map_sheet.max_row + 1):  
        map_sheet.cell(  
            row=row,  
            column=7,  
            value=f'=B{row}&"|"&D{row}'  
        )  

        map_sheet.cell(  
            row=row,  
            column=8,  
            value=f"=F{row}"  
        )  

        map_sheet.cell(  
            row=row,  
            column=9,  
            value=f'=B{row}&"|"&E{row}'  
        )  

        map_sheet.cell(  
            row=row,  
            column=10,  
            value=f"=F{row}"  
        )  

    headers = {  
        validation_sheet.cell(  
            row=1,  
            column=column  
        ).value: column  
        for column in range(  
            1,  
            validation_sheet.max_column + 1  
        )  
    }  

    polygon_type_col = headers["polygon_type"]  
    kabupaten_col = headers["kabupaten"]  
    wilayah_col = headers["wilayah"]  

    lookup_key_col = validation_sheet.max_column + 1  
    formula_id_col = lookup_key_col + 1  

    validation_sheet.cell(  
        row=1,  
        column=lookup_key_col,  
        value="LOOKUP_KEY"  
    )  

    validation_sheet.cell(  
        row=1,  
        column=formula_id_col,  
        value="ID_AREA_VLOOKUP"  
    )  

    data_last_row = len(  
        list(validation_sheet.iter_rows())  
    ) + 1  

    for row in range(2, data_last_row):  
        type_letter = validation_sheet.cell(  
            row=row,  
            column=polygon_type_col  
        ).column_letter  

        lookup_key_letter = validation_sheet.cell(  
            row=row,  
            column=lookup_key_col  
        ).column_letter  

        validation_sheet.cell(  
            row=row,  
            column=lookup_key_col,  
            value=(  
                f'=UPPER('  
                f'{validation_sheet.cell(row=row, column=kabupaten_col).column_letter}{row}'  
                f'&"|"&'  
                f'{validation_sheet.cell(row=row, column=wilayah_col).column_letter}{row}'  
                f')'  
            )  
        )  

        validation_sheet.cell(  
            row=row,  
            column=formula_id_col,  
            value=(  
                f'=IF({type_letter}{row}="Grid","SKIPPED",'  
                f'IF({type_letter}{row}="Desa",'  
                f'IFERROR(VLOOKUP('  
                f'{lookup_key_letter}{row},'  
                f"'Map Border'!$G:$H,2,FALSE),"  
                f'"NOT FOUND"),'  
                f'IF({type_letter}{row}="City",'  
                f'IFERROR(VLOOKUP('  
                f'{lookup_key_letter}{row},'  
                f"'Map Border'!$I:$J,2,FALSE),"  
                f'"NOT FOUND"),'  
                f'"INVALID TYPE")))'  
            )  
        )  

    header_fill = PatternFill(  
        fill_type="solid",  
        fgColor="1F4E78"  
    )  

    header_font = Font(  
        bold=True,  
        color="FFFFFF"  
    )  

    for sheet in [validation_sheet, map_sheet]:  
        for cell in sheet[1]:  
            cell.fill = header_fill  
            cell.font = header_font  
            cell.alignment = Alignment(  
                horizontal="center",  
                vertical="center"  
            )  

        for column_cells in sheet.columns:  
            max_length = 0  

            column_letter = get_column_letter(  
                column_cells[0].column  
            )  

            for cell in column_cells:  
                max_length = max(  
                    max_length,  
                    len(str(cell.value))  
                )  

            sheet.column_dimensions[  
                column_letter  
            ].width = min(max_length + 2, 50)  

    workbook.calculation.fullCalcOnLoad = True  
    workbook.calculation.forceFullCalc = True  
    workbook.calculation.calcMode = "auto"  

    workbook.save(output_file)  

# ======================================  
# PROCESS AND VALIDATE  
# ======================================  

def process_and_validate_file(  
    file_path,  
    mode,  
    output_folder  
):  
    try:  
        df = pd.read_csv(  
            file_path,  
            dtype=str,  
            keep_default_na=False  
        )  

    except Exception as e:  
        return (  
            None,  
            None,  
            None,  
            f"❌ Gagal membaca file:\n\n{str(e)}"  
        )  

    try:  
        map_df = load_map_border()  

    except Exception as e:  
        return (  
            None,  
            None,  
            None,  
            f"❌ Gagal memuat Map Border:\n\n{str(e)}"  
        )  

    df.columns = (  
        df.columns  
        .str.strip()  
        .str.lower()  
    )  

    expected_columns = (  
        UPDATE_COLUMNS  
        if mode == "update"  
        else BASE_COLUMNS  
    )  

    missing_columns = [  
        col  
        for col in expected_columns  
        if col not in df.columns  
    ]  

    if missing_columns:  
        report = (  
            "❌ Header tidak sesuai\n\n"  
            f"Mode : {mode.upper()}\n"  
            f"Kolom kurang:\n{missing_columns}\n\n"  
            f"Header ditemukan:\n{list(df.columns)}"  
        )  

        return (  
            None,  
            None,  
            None,  
            report  
        )  

    df = df[expected_columns].copy()  

    df["_original_row_number"] = range(  
        2,  
        len(df) + 2  
    )  

    for col in expected_columns:  
        df[col] = (  
            df[col]  
            .fillna("")  
            .astype(str)  
            .str.strip()  
        )  

    df["_polygon_type_upper"] = (  
        df["polygon_type"]  
        .str.upper()  
    )  

    df["polygon_type"] = (  
        df["polygon_type"]  
        .apply(format_title_case)  
    )  

    total_rows = len(df)  

    df_unique = (  
        df  
        .drop_duplicates(  
            subset=expected_columns,  
            keep="first"  
        )  
        .reset_index(drop=True)  
    )  

    duplicate_count = total_rows - len(df_unique)  

    map_border_result = create_map_border_result(  
        df_unique,  
        map_df  
    )  

    polygon_label_check = (  
        df_unique  
        .groupby("polygon")["label"]  
        .nunique()  
    )  

    conflict_polygon = (  
        polygon_label_check[  
            polygon_label_check > 1  
        ]  
        .index  
        .tolist()  
    )  

    content_errors = validate_content(  
        df_unique,  
        expected_columns,  
        map_df  
    )  

    for polygon in conflict_polygon:  
        conflict_rows = df_unique[  
            df_unique["polygon"] == polygon  
        ]  

        labels = (  
            conflict_rows["label"]  
            .unique()  
            .tolist()  
        )  

        for _, row in conflict_rows.iterrows():  
            content_errors.append({  
                "row_number": row["_original_row_number"],  
                "polygon": polygon,  
                "error": (  
                    "Polygon memiliki lebih dari satu label: "  
                    + ", ".join(labels)  
                )  
            })  

    if content_errors:  
        status = "❌ FAILED"  
    elif duplicate_count > 0:  
        status = "⚠️ PASS WITH WARNING"  
    else:  
        status = "✅ PASS"  

    report = "📊 SIFA VALIDATION REPORT\n\n"  
    report += f"MODE : {mode.upper()}\n"  
    report += f"STATUS : {status}\n\n"  
    report += f"📌 Total Input : {total_rows}\n"  
    report += f"✨ Unique Data : {len(df_unique)}\n"  
    report += f"🗑 Duplicate   : {duplicate_count}\n\n"  

    label_count = df_unique["label"].value_counts()  

    report += "📋 LIST UNIQUE LABELS & TOTAL:\n"  

    for label, count in label_count.items():  
        report += f"▫️ {label}: {count} baris\n"  

    report += "\n"  

    if content_errors:  
        report += (  
            f"⚠️ Total Error : {len(content_errors)}\n"  
            "Detail error dikirim dalam file terpisah.\n\n"  
            "Note : Silakan divalidasi ulang untuk memastikan "  
            "kelancaran Activity Changes.\n\n"  
            "❌ File memiliki error dan tidak boleh langsung "  
            "digunakan sebelum diperbaiki.\n\n"  
        )  

    if duplicate_count > 0:  
        report += (  
            "⚠️ Duplicate ditemukan dan sudah dihapus "  
            "dari file output.\n\n"  
        )  

    if mode == "update":  
        output_df = df_unique[  
            [  
                "polygon",  
                "polygon_type",  
                "label",  
                "label_new"  
            ]  
        ].copy()  

        output_df["start_date"] = ""  
        output_df["end_date"] = ""  

    else:  
        output_df = df_unique[  
            [  
                "polygon",  
                "polygon_type",  
                "label"  
            ]  
        ].copy()  

    filename = os.path.splitext(  
        os.path.basename(file_path)  
    )[0]  

    if mode == "delete":  
        output_file = os.path.join(  
            output_folder,  
            f"clean_delete_{filename}.txt"  
        )  

        output_df.to_csv(  
            output_file,  
            sep="|",  
            index=False  
        )  

        report += (  
            "📤 Output : DELETE\n"  
            "Separator : |\n"  
        )  

    elif mode == "update":  
        output_file = os.path.join(  
            output_folder,  
            f"clean_update_{filename}.txt"  
        )  

        output_df.to_csv(  
            output_file,  
            sep=",",  
            index=False  
        )  

        report += (  
            "📤 Output : UPDATE\n"  
            "Separator : ,\n"  
        )  

    else:  
        output_file = os.path.join(  
            output_folder,  
            f"clean_insert_{filename}.txt"  
        )  

        output_df.to_csv(  
            output_file,  
            sep=",",  
            index=False  
        )  

        report += (  
            "📤 Output : INSERT\n"  
            "Separator : ,\n"  
        )  

    error_file = None  

    if content_errors:  
        error_file = os.path.join(  
            output_folder,  
            f"validation_error_{filename}.csv"  
        )  

        error_df = pd.DataFrame(  
            content_errors,  
            columns=[  
                "row_number",  
                "polygon",  
                "error"  
            ]  
        )  

        error_df.to_csv(  
            error_file,  
            index=False,  
            encoding="utf-8-sig"  
        )  

    map_border_file = os.path.join(  
        output_folder,  
        f"map_border_validation_{filename}.xlsx"  
    )  

    create_map_border_excel(  
        input_df=df_unique,  
        map_border_result=map_border_result,  
        map_df=map_df,  
        output_file=map_border_file,  
        mode=mode  
    )  

    add_excel_vlookup_formulas(  
        output_file=map_border_file  
    )  

    return (  
        output_file,  
        error_file,  
        map_border_file,  
        report  
    )  

# ======================================  
# PROGRESS MESSAGE  
# ======================================  

async def send_progress(update: Update, start_time: float):  
    messages = [  
        "⏳ Masih diproses, jangan bosan ya wkwk...",  
        "🔄 Validasi CSV masih berjalan...",  
        "📊 Sedang membuat laporan Excel...",  
        "🗺️ Sedang memproses Map Border...",  
        "⌛ Proses masih berjalan, mohon tunggu..."  
    ]  

    index = 0  
    loop = asyncio.get_running_loop()  

    try:  
        while True:  
            await asyncio.sleep(PROGRESS_INTERVAL)  

            elapsed_minutes = int(  
                (loop.time() - start_time) / 60  
            )  

            await update.message.reply_text(  
                f"{messages[index % len(messages)]}\n"  
                f"⏱️ Waktu berjalan: {elapsed_minutes} menit"  
            )  

            index += 1  

    except asyncio.CancelledError:  
        pass  

# ======================================  
# START COMMAND  
# ======================================  

async def start_command(  
    update: Update,  
    context: ContextTypes.DEFAULT_TYPE  
):  
    if not update.message:  
        return  

    await update.message.reply_text(  
        "<b>TOP GLOBAL SIFA is HERE !!</b>\n\n"  
        "Cara penggunaan:\n\n"  
        "/val_sifa insert\n"  
        "/val_sifa delete\n"  
        "/val_sifa update\n\n"  
        "<b>Upload CSV + header dengan caption command.</b>\n\n"  
        "<b>- File Insert dan Delete harus berisi kolom:</b>\n"  
        "polygon, polygon_type, label\n\n"  
        "<b>- File Update harus berisi kolom:</b>\n"  
        "polygon,polygon_type,label,label_new\n\n"  
        "<b>Validasi polygon:</b>\n"  
        "- DESA: KABUPATEN_DESA\n"  
        "- CITY: KABUPATEN_AREA\n"  
        "- GRID: tidak divalidasi ke Map Border\n\n",  
        parse_mode="HTML"  
    )  

# ======================================  
# FILE HANDLER  
# ======================================  

async def handle_document(  
    update: Update,  
    context: ContextTypes.DEFAULT_TYPE  
):  
    if not update.message or not update.message.document:  
        return  

    caption = update.message.caption  

    if not caption:  
        return  

    command = caption.lower().split()  
    cmd = command[0].split("@")[0]
    if len(command) != 2:  
        cmd = command[0].split("@")[0]  

    if cmd != "/val_sifa":  
        await update.message.reply_text(  
            "❌ Format salah\n\n"  
            "Gunakan:\n"  
            "/val_sifa insert\n"  
            "/val_sifa delete\n"  
            "/val_sifa update"  
        )  
        return  

    mode = command[1]  

    if mode not in ["insert", "delete", "update"]:  
        await update.message.reply_text(  
            "❌ Mode tidak tersedia\n\n"  
            "Gunakan:\n"  
            "/val_sifa insert\n"  
            "/val_sifa delete\n"  
            "/val_sifa update"  
        )  
        return  

    document = update.message.document  
    filename = document.file_name or "uploaded.csv"  

    if not filename.lower().endswith(".csv"):  
        await update.message.reply_text(  
            "❌ File harus berformat CSV"  
        )  
        return  

    await update.message.reply_text(  
        f"⏳ Processing {filename}\n"  
        "Sabar nggih, will update every 5 minutes kok, maksimal proses sejam :v"  
    )  

    try:  
        with tempfile.TemporaryDirectory() as temp_folder:  
            local_file = os.path.join(  
                temp_folder,  
                filename  
            )  

            telegram_file = await context.bot.get_file(  
                document.file_id,  
                read_timeout=TELEGRAM_TIMEOUT,  
                connect_timeout=60.0,  
                pool_timeout=60.0  
            )  

            await telegram_file.download_to_drive(  
                local_file,  
                read_timeout=TELEGRAM_TIMEOUT,  
                write_timeout=TELEGRAM_TIMEOUT,  
                connect_timeout=60.0,  
                pool_timeout=60.0  
            )  

            start_time = asyncio.get_running_loop().time()  

            progress_task = asyncio.create_task(  
                send_progress(update, start_time)  
            )  

            try:  
                loop = asyncio.get_running_loop()  

                result = await loop.run_in_executor(  
                    None,  
                    process_and_validate_file,  
                    local_file,  
                    mode,  
                    temp_folder  
                )  

                (  
                    output_file,  
                    error_file,  
                    map_border_file,  
                    report  
                ) = result  

            finally:  
                progress_task.cancel()  

                try:  
                    await progress_task  
                except asyncio.CancelledError:  
                    pass  

            await update.message.reply_text(report)  

            if output_file:  
                with open(output_file, "rb") as file:  
                    await update.message.reply_document(  
                        document=file,  
                        filename=os.path.basename(output_file),  
                        caption=(  
                            "✅ File hasil validasi "  
                            "dan cleaning"  
                        ),  
                        read_timeout=TELEGRAM_TIMEOUT,  
                        write_timeout=TELEGRAM_TIMEOUT,  
                        connect_timeout=60.0,  
                        pool_timeout=60.0  
                    )  

            if map_border_file:  
                with open(map_border_file, "rb") as file:  
                    await update.message.reply_document(  
                        document=file,  
                        filename=os.path.basename(map_border_file),  
                        caption=(  
                            "🗺️ Hasil validasi Map Border"  
                        ),  
                        read_timeout=TELEGRAM_TIMEOUT,  
                        write_timeout=TELEGRAM_TIMEOUT,  
                        connect_timeout=60.0,  
                        pool_timeout=60.0  
                    )  

            if error_file:  
                with open(error_file, "rb") as file:  
                    await update.message.reply_document(  
                        document=file,  
                        filename=os.path.basename(error_file),  
                        caption=(  
                            "⚠️ Detail baris yang "  
                            "memiliki error"  
                        ),  
                        read_timeout=TELEGRAM_TIMEOUT,  
                        write_timeout=TELEGRAM_TIMEOUT,  
                        connect_timeout=60.0,  
                        pool_timeout=60.0  
                    )  

    except Exception as e:  
        logging.exception("System error")  

        await update.message.reply_text(  
            f"❌ Error sistem:\n{str(e)}"  
        )  

# ======================================  
# MAIN  
# ======================================  

def main():  
    request = HTTPXRequest(  
        connection_pool_size=20,  
        connect_timeout=60.0,  
        read_timeout=TELEGRAM_TIMEOUT,  
        write_timeout=TELEGRAM_TIMEOUT,  
        pool_timeout=60.0  
    )  

    app = (  
        Application  
        .builder()  
        .token(TOKEN)  
        .request(request)  
        .get_updates_request(request)  
        .build()  
    )  

    app.add_handler(  
        CommandHandler(  
            "sifa",  
            start_command  
        )  
    )  

    app.add_handler(  
        MessageHandler(  
            filters.Document.ALL,  
            handle_document,  
            block=False  
        )  
    )  

    print("🤖 SIFA Validator Bot Running...")  

    app.run_polling()  

if __name__ == "__main__":  
    main()  